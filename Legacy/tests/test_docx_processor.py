"""
Unit tests for DOCX processing in src/docx_processor.py and its wiring into
src/agent.py. No Ollama needed: _extract_with_llm is monkeypatched.
"""

import io
import sys
from pathlib import Path

import docx

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src.docx_processor import DocxProcessor, DocxValidator
from src.agent import OrderExtractionPipeline


def make_docx_bytes() -> io.BytesIO:
    document = docx.Document()
    document.add_paragraph("Order from Jane Doe")
    document.add_paragraph("Ship to: 123 Main St, Springfield, IL 62704")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[1].text = "Qty"
    table.rows[1].cells[0].text = "Widget"
    table.rows[1].cells[1].text = "2"

    buf = io.BytesIO()
    document.save(buf)
    buf.seek(0)
    return buf


def test_process_docx_extracts_text_and_tables():
    result = DocxProcessor(chunk_size=2000).process_docx(make_docx_bytes())

    assert result["success"] is True
    assert "Order from Jane Doe" in result["text"]
    assert result["tables"] == [[["Item", "Qty"], ["Widget", "2"]]]
    assert result["num_chunks"] == 1


def test_process_docx_reports_error_on_invalid_file():
    result = DocxProcessor().process_docx(io.BytesIO(b"not a docx"))
    assert result["success"] is False
    assert "error" in result


def test_docx_validator_checks_header_and_size():
    valid_bytes = make_docx_bytes()
    assert DocxValidator.is_valid_docx(valid_bytes) == (True, "")
    assert DocxValidator.is_valid_docx(io.BytesIO(b"not a docx"))[0] is False

    is_valid, _ = DocxValidator.check_file_size(valid_bytes, max_size_mb=0.0001)
    assert is_valid is False


def test_process_docx_streaming_routes_through_shared_pipeline():
    agent = OrderExtractionPipeline()
    agent._extract_with_llm = lambda text: {"customer_name": "Jane Doe", "items": []}

    result = agent.process_docx(make_docx_bytes())

    assert result["order"]["customer_name"] == "Jane Doe"
    assert result["extraction_metadata"]["source_type"] == "docx"
    assert result["extraction_metadata"]["docx_info"]["num_tables"] == 1


if __name__ == "__main__":
    import pytest
    raise SystemExit(pytest.main([__file__, "-v"]))
