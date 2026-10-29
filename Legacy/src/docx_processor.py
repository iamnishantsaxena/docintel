"""
DOCX Processing Module
Handles Word document parsing, text/table extraction, and chunking
(reuses the text-cleaning/chunking machinery in pdf_processor.py)
"""

import logging
from typing import Dict, List, Tuple
from pathlib import Path

try:
    import docx
except ImportError:
    raise ImportError("Please install python-docx: pip install python-docx")

from .pdf_processor import DocumentProcessor, check_file_size, is_valid_zip_container

logger = logging.getLogger(__name__)


class DocxProcessor(DocumentProcessor):
    """
    DOCX processing: paragraph + table text extraction, then shared chunking.
    """

    def process_docx(self, docx_file) -> Dict[str, any]:
        """
        Main DOCX processing method

        Args:
            docx_file: File path or file-like object

        Returns:
            Dictionary with extracted text, metadata, chunks, tables (same shape as
            PDFProcessor.process_pdf, minus page count — DOCX has no fixed pagination)
        """
        try:
            text, metadata, tables = self._extract_docx(docx_file)
            text = self._clean_text(text)
            chunks = self._create_chunks(text)

            return {
                "text": text,
                "metadata": metadata,
                "chunks": chunks,
                "tables": tables,
                "total_length": len(text),
                "num_chunks": len(chunks),
                "success": True
            }

        except Exception as e:
            logger.error(f"Error processing DOCX: {e}")
            return {
                "text": "",
                "metadata": {},
                "chunks": [],
                "tables": [],
                "error": str(e),
                "success": False
            }

    def _extract_docx(self, docx_file) -> Tuple[str, Dict, List[List[List[str]]]]:
        """Extract paragraph text, core metadata, and tables from a .docx file"""
        if isinstance(docx_file, (str, Path)):
            document = docx.Document(str(docx_file))
        else:
            docx_file.seek(0)
            document = docx.Document(docx_file)

        text_parts = [p.text for p in document.paragraphs if p.text.strip()]

        props = document.core_properties
        metadata = {
            "num_paragraphs": len(document.paragraphs),
            "extractor": "python-docx",
            "title": props.title or "",
            "author": props.author or "",
            "subject": props.subject or "",
        }

        tables = [
            [[cell.text for cell in row.cells] for row in table.rows]
            for table in document.tables
        ]

        return "\n".join(text_parts), metadata, tables

    def get_summary(self, docx_result: Dict) -> str:
        """Get a summary of the DOCX processing results"""
        if not docx_result.get("success"):
            return f"DOCX processing failed: {docx_result.get('error', 'Unknown error')}"

        summary_parts = [
            f"Paragraphs: {docx_result['metadata'].get('num_paragraphs', 'unknown')}",
            f"Text length: {docx_result['total_length']} characters",
            f"Chunks: {docx_result['num_chunks']}",
            f"Tables: {len(docx_result['tables'])}",
        ]

        return " | ".join(summary_parts)


class DocxValidator:
    """Validate DOCX files before processing"""

    is_valid_docx = staticmethod(is_valid_zip_container)
    check_file_size = staticmethod(check_file_size)
