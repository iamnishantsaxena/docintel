"""
Excel Processing Module
Handles .xlsx parsing: detects the order-items header row on the first sheet,
splits it from any scalar front-matter (customer/address key-value rows) above it.
"""

import logging
from typing import Dict, List, Optional, Tuple
from pathlib import Path

try:
    import openpyxl
except ImportError:
    raise ImportError("Please install openpyxl: pip install openpyxl")

from .pdf_processor import DocumentProcessor, check_file_size, is_valid_zip_container

logger = logging.getLogger(__name__)

# A row counts as the order-items header row once at least two cells match these
# ponytail: keyword list, not a real header-detection model — extend if real sheets miss it
HEADER_KEYWORDS = {
    "item", "product", "description", "qty", "quantity", "price", "unit price",
    "amount", "sku", "total", "name", "cost", "subtotal",
}


class ExcelProcessor(DocumentProcessor):
    """
    Excel processing: front-matter (key/value rows) + order-items table on the
    first sheet, then shared chunking/table-formatting from DocumentProcessor.
    """

    def process_excel(self, excel_file) -> Dict[str, any]:
        """
        Main Excel processing method

        Args:
            excel_file: File path or file-like object

        Returns:
            Dictionary with extracted text, metadata, chunks, tables (same shape as
            PDFProcessor.process_pdf / DocxProcessor.process_docx)
        """
        try:
            text, metadata, tables = self._extract_excel(excel_file)
            text = self._clean_text(text)
            chunks = self._create_chunks(text)

            return {
                "text": text,
                "metadata": metadata,
                "chunks": chunks,
                "tables": tables,
                "total_length": len(text),
                "num_chunks": len(chunks),
                "success": True
            }

        except Exception as e:
            logger.error(f"Error processing Excel file: {e}")
            return {
                "text": "",
                "metadata": {},
                "chunks": [],
                "tables": [],
                "error": str(e),
                "success": False
            }

    def _extract_excel(self, excel_file) -> Tuple[str, Dict, List[List[List[str]]]]:
        """Split the first sheet into front-matter text and an order-items table"""
        if isinstance(excel_file, (str, Path)):
            workbook = openpyxl.load_workbook(str(excel_file), data_only=True, read_only=True)
        else:
            excel_file.seek(0)
            workbook = openpyxl.load_workbook(excel_file, data_only=True, read_only=True)

        # ponytail: first/active sheet only, multi-sheet merge if a real workbook needs it
        sheet = workbook.active
        rows = [
            ["" if cell is None else str(cell) for cell in row]
            for row in sheet.iter_rows(values_only=True)
        ]
        # Drop fully-blank rows (common in exported spreadsheets)
        rows = [row for row in rows if any(cell.strip() for cell in row)]

        metadata = {
            "num_rows": len(rows),
            "num_sheets": len(workbook.sheetnames),
            "sheet_name": sheet.title,
            "extractor": "openpyxl",
        }

        header_idx = self._find_header_row(rows)
        if header_idx is None:
            front_matter, table_rows = rows, []
        else:
            front_matter, table_rows = rows[:header_idx], rows[header_idx:]

        text = "\n".join(self._row_to_line(row) for row in front_matter)
        tables = [table_rows] if table_rows else []

        return text, metadata, tables

    def _find_header_row(self, rows: List[List[str]]) -> Optional[int]:
        """First row with >=2 cells matching a known order-items header keyword"""
        for idx, row in enumerate(rows):
            matches = sum(1 for cell in row if cell.strip().lower() in HEADER_KEYWORDS)
            if matches >= 2:
                return idx
        return None

    def _row_to_line(self, row: List[str]) -> str:
        """Render a front-matter row as text: 'Key: Value' for a two-cell row, else space-joined"""
        cells = [c for c in row if c.strip()]
        if len(cells) == 2:
            return f"{cells[0]}: {cells[1]}"
        return " ".join(cells)

    def get_summary(self, excel_result: Dict) -> str:
        """Get a summary of the Excel processing results"""
        if not excel_result.get("success"):
            return f"Excel processing failed: {excel_result.get('error', 'Unknown error')}"

        summary_parts = [
            f"Rows: {excel_result['metadata'].get('num_rows', 'unknown')}",
            f"Text length: {excel_result['total_length']} characters",
            f"Chunks: {excel_result['num_chunks']}",
            f"Tables: {len(excel_result['tables'])}",
        ]

        return " | ".join(summary_parts)


class ExcelValidator:
    """Validate Excel files before processing"""

    is_valid_excel = staticmethod(is_valid_zip_container)
    check_file_size = staticmethod(check_file_size)
