import csv
import hashlib
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import worker
from app.db import Document, SessionLocal, Status, User
from app.extraction import LLMUnavailable
from app.main import app
from eval.pdfgen import make_pdf
from tests.conftest import GOOD_LLM, INVOICE_LINES


@pytest.fixture
def client():
    with TestClient(app) as c:
        yield c


def auth(client, email="alice@example.com"):
    r = client.post("/api/auth/register", json={"email": email, "password": "correct-horse"})
    assert r.status_code == 201
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def upload(client, headers, *files):
    return client.post("/api/documents", headers=headers,
                       files=[("files", (name, data, "application/pdf")) for name, data in files])


# ---------- auth ----------

def test_register_login_me(client):
    auth(client)
    r = client.post("/api/auth/login", data={"username": "ALICE@example.com", "password": "correct-horse"})
    token = r.json()["access_token"]
    assert client.get("/api/auth/me", headers={"Authorization": f"Bearer {token}"}).json()["email"] == "alice@example.com"


def test_bad_credentials_and_duplicates(client):
    auth(client)
    assert client.post("/api/auth/login", data={"username": "alice@example.com", "password": "wrong-pass"}).status_code == 401
    assert client.post("/api/auth/register", json={"email": "alice@example.com", "password": "another-pw"}).status_code == 409
    assert client.post("/api/auth/register", json={"email": "not-an-email", "password": "x" * 8}).status_code == 422
    assert client.post("/api/auth/register", json={"email": "b@example.com", "password": "short"}).status_code == 422


@pytest.mark.parametrize("header", [None, "Bearer garbage", "Bearer eyJhbGciOiJIUzI1NiJ9.e30.x"])
def test_protected_routes_require_valid_token(client, header):
    headers = {"Authorization": header} if header else {}
    assert client.get("/api/documents", headers=headers).status_code == 401


def test_users_cannot_see_each_others_documents(client, llm, invoice_pdf):
    alice = auth(client)
    doc_id = upload(client, alice, ("a.pdf", invoice_pdf)).json()[0]["id"]
    bob = auth(client, "bob@example.com")
    for method, path in [("get", f"/api/documents/{doc_id}"), ("get", f"/api/documents/{doc_id}/file"),
                         ("post", f"/api/documents/{doc_id}/approve"), ("delete", f"/api/documents/{doc_id}")]:
        assert getattr(client, method)(path, headers=bob).status_code == 404
    assert client.get("/api/documents", headers=bob).json() == []
    # Same file uploaded by another user is that user's own, separate document.
    assert upload(client, bob, ("a.pdf", invoice_pdf)).json()[0]["duplicate"] is False


# ---------- upload validation ----------

def test_upload_rejects_non_pdfs_and_oversize(client, llm, monkeypatch):
    h = auth(client)
    monkeypatch.setattr(worker.settings, "max_upload_mb", 0.001)
    res = upload(client, h, ("evil.pdf", b"<script>alert(1)</script>"), ("empty.pdf", b""),
                 ("big.pdf", b"%PDF-" + b"x" * 5000))
    assert [r["error"] for r in res.json()] == ["Only PDF files are accepted", "Empty file", "File exceeds size limit"]
    assert client.get("/api/documents", headers=h).json() == []


def test_path_traversal_filename_is_neutralised(client, llm, invoice_pdf):
    h = auth(client)
    res = upload(client, h, ("../../etc/passwd.pdf", invoice_pdf)).json()[0]
    assert res["filename"] == "passwd.pdf"


def test_malformed_pdf_fails_cleanly(client, llm):
    h = auth(client)
    doc = upload(client, h, ("broken.pdf", b"%PDF-1.4 not really")).json()[0]
    detail = client.get(f"/api/documents/{doc['id']}", headers=h).json()
    assert detail["status"] == "failed" and "Could not parse" in detail["error"]


# ---------- full flow ----------

def test_upload_review_approve_export(client, llm, invoice_pdf):
    h = auth(client)
    res = upload(client, h, ("inv.pdf", invoice_pdf), ("po.pdf", make_pdf([["PURCHASE ORDER", "PO Number: PO-77"]])))
    assert res.status_code == 202 and len(res.json()) == 2
    doc_id = res.json()[0]["id"]

    detail = client.get(f"/api/documents/{doc_id}", headers=h).json()
    assert detail["status"] == "needs_review" and detail["fields"]["total"] == "495.00"
    assert detail["issues"] == [] and detail["page_count"] == 1
    assert client.get(f"/api/documents/{doc_id}/file", headers=h).content.startswith(b"%PDF")

    # Reviewer introduces an error -> approval blocked server-side
    bad = {**detail["fields"], "total": "500.00"}
    r = client.put(f"/api/documents/{doc_id}/fields", headers=h, json=bad)
    assert any(i["code"] == "total_mismatch" for i in r.json()["issues"])
    assert client.post(f"/api/documents/{doc_id}/approve", headers=h).status_code == 422

    # ... fixes it, approves
    r = client.put(f"/api/documents/{doc_id}/fields", headers=h, json={**bad, "total": "495.00", "supplier_name": "=HYPERLINK(1)"})
    assert r.json()["issues"] == []
    assert client.post(f"/api/documents/{doc_id}/approve", headers=h).json()["status"] == "validated"
    assert client.get(f"/api/documents/{doc_id}", headers=h).json()["extracted"]["supplier_name"] == "ACME Supplies Ltd"

    rows = list(csv.DictReader(io.StringIO(client.get("/api/export?format=csv", headers=h).text)))
    assert len(rows) == 1 and rows[0]["total"] == "495.00"
    assert rows[0]["supplier_name"] == "'=HYPERLINK(1)"  # formula injection neutralised

    wb = load_workbook(io.BytesIO(client.get("/api/export?format=xlsx", headers=h).content))
    assert wb["documents"].max_row == 2 and wb["line_items"].max_row == 3
    assert float(wb["documents"]["M2"].value) == 495.0


def test_put_fields_validates_input(client, llm, invoice_pdf):
    h = auth(client)
    doc_id = upload(client, h, ("inv.pdf", invoice_pdf)).json()[0]["id"]
    for body in [{"total": "lots"}, {"document_date": "2024-13-45"}, {"unknown": 1}, {"currency": "DOLLARS"}]:
        assert client.put(f"/api/documents/{doc_id}/fields", headers=h, json=body).status_code == 422


# ---------- duplicates, idempotency, retries ----------

def test_duplicate_upload_returns_existing_document(client, llm, invoice_pdf):
    h = auth(client)
    first = upload(client, h, ("a.pdf", invoice_pdf)).json()[0]
    second = upload(client, h, ("renamed.pdf", invoice_pdf)).json()[0]
    assert second == {**second, "id": first["id"], "duplicate": True}
    assert llm.calls == 1


def test_same_invoice_number_different_file_flagged(client, llm, invoice_pdf):
    h = auth(client)
    upload(client, h, ("a.pdf", invoice_pdf))
    other = make_pdf([[*INVOICE_LINES, "Scanned copy"]])
    doc_id = upload(client, h, ("b.pdf", other)).json()[0]["id"]
    issues = client.get(f"/api/documents/{doc_id}", headers=h).json()["issues"]
    assert [i["code"] for i in issues] == ["possible_duplicate"]


def test_task_is_idempotent(client, llm, invoice_pdf):
    h = auth(client)
    doc_id = upload(client, h, ("a.pdf", invoice_pdf)).json()[0]["id"]
    assert worker.run(doc_id) == "skipped"  # already processed; a redelivered task is a no-op
    assert llm.calls == 1


def _queued_doc(invoice_pdf):
    sha = hashlib.sha256(invoice_pdf).hexdigest()
    worker.store_file(sha, invoice_pdf)
    with SessionLocal() as db:
        user = User(email="w@example.com", password_hash="x")
        db.add(user)
        db.flush()
        doc = Document(owner_id=user.id, filename="a.pdf", sha256=sha, size_bytes=1, status=Status.queued)
        db.add(doc)
        db.commit()
        return doc.id


def test_llm_outage_requeues_then_degrades_to_regex(llm, invoice_pdf):
    doc_id = _queued_doc(invoice_pdf)
    llm.reply = LLMUnavailable("timeout")
    with pytest.raises(worker.Retryable):
        worker.run(doc_id)
    with SessionLocal() as db:
        assert db.get(Document, doc_id).status == Status.queued  # back in the queue for the retry

    assert worker.run(doc_id, final_attempt=True) == "done"
    with SessionLocal() as db:
        doc = db.get(Document, doc_id)
        assert doc.status == Status.needs_review and doc.attempts == 2
        assert doc.fields["total"] == "495.00"
        assert any(i["code"] == "llm_unavailable" for i in doc.issues)


def test_celery_task_retries_then_succeeds(llm, invoice_pdf, monkeypatch):
    doc_id = _queued_doc(invoice_pdf)
    replies = iter([LLMUnavailable("timeout"), None])

    def flaky(text):
        r = next(replies)
        if r:
            raise r
        return dict(GOOD_LLM)
    monkeypatch.setattr("app.extraction.call_llm", flaky)
    worker.process_document.apply(args=[doc_id])  # eager: retries run inline
    with SessionLocal() as db:
        doc = db.get(Document, doc_id)
        assert doc.status == Status.needs_review and doc.attempts == 2


def test_reprocess_failed_document(client, llm):
    h = auth(client)
    doc_id = upload(client, h, ("broken.pdf", b"%PDF-1.4 nope")).json()[0]["id"]
    r = client.post(f"/api/documents/{doc_id}/reprocess", headers=h)
    assert r.status_code == 202 and r.json()["status"] == "failed"  # still broken, but re-ran
    assert r.json()["attempts"] == 2
