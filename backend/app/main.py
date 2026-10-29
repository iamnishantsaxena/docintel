import csv
import hashlib
import io
from contextlib import asynccontextmanager
from datetime import datetime
from decimal import Decimal
from typing import Literal

from fastapi import Depends, FastAPI, HTTPException, Query, UploadFile, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response, StreamingResponse
from fastapi.security import OAuth2PasswordRequestForm
from openpyxl import Workbook
from pydantic import BaseModel, EmailStr, Field
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from . import rules
from .auth import create_token, current_user, get_db, hash_password, verify_password
from .config import settings
from .db import Document, Status, User, init_db, utcnow
from .fields import Fields
from .worker import enqueue, file_path, store_file, validate_document


@asynccontextmanager
async def lifespan(_app):
    init_db()
    yield


app = FastAPI(title="Document Intelligence API", version="2.0.0", lifespan=lifespan,
              docs_url="/api/docs", openapi_url="/api/openapi.json")
app.add_middleware(CORSMiddleware, allow_origins=settings.cors_origins, allow_methods=["*"],
                   allow_headers=["Authorization", "Content-Type"])


# ---------- schemas ----------

class Credentials(BaseModel):
    email: EmailStr
    password: str = Field(min_length=8, max_length=128)


class Token(BaseModel):
    access_token: str
    token_type: str = "bearer"


class Issue(BaseModel):
    field: str | None
    code: str
    severity: Literal["error", "warning"]
    message: str


class DocumentSummary(BaseModel):
    id: int
    filename: str
    status: Status
    error: str | None
    error_count: int
    warning_count: int
    supplier_name: str | None
    document_number: str | None
    total: str | None
    created_at: datetime
    updated_at: datetime


class DocumentDetail(DocumentSummary):
    page_count: int | None
    attempts: int
    fields: Fields | None
    extracted: Fields | None
    issues: list[Issue]
    meta: dict | None
    reviewed_at: datetime | None


class UploadResult(BaseModel):
    id: int | None
    filename: str
    status: Status | None
    duplicate: bool = False
    error: str | None = None


def summary(doc: Document) -> dict:
    issues = doc.issues or []
    f = doc.fields or {}
    return {
        "id": doc.id, "filename": doc.filename, "status": doc.status, "error": doc.error,
        "error_count": sum(i["severity"] == "error" for i in issues),
        "warning_count": sum(i["severity"] == "warning" for i in issues),
        "supplier_name": f.get("supplier_name"), "document_number": f.get("document_number"),
        "total": f.get("total"), "created_at": doc.created_at, "updated_at": doc.updated_at,
    }


def detail(doc: Document) -> dict:
    return {**summary(doc), "page_count": doc.page_count, "attempts": doc.attempts, "fields": doc.fields,
            "extracted": doc.extracted, "issues": doc.issues or [], "meta": doc.meta, "reviewed_at": doc.reviewed_at}


def owned_document(doc_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)) -> Document:
    doc = db.get(Document, doc_id)
    # 404 rather than 403 for other users' documents: don't reveal that the id exists.
    if not doc or doc.owner_id != user.id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Document not found")
    return doc


# ---------- auth ----------

@app.post("/api/auth/register", response_model=Token, status_code=201, tags=["auth"])
def register(body: Credentials, db: Session = Depends(get_db)):
    user = User(email=body.email.lower(), password_hash=hash_password(body.password))
    db.add(user)
    try:
        db.commit()
    except IntegrityError:
        raise HTTPException(status.HTTP_409_CONFLICT, "Email already registered") from None
    return Token(access_token=create_token(user.id))


@app.post("/api/auth/login", response_model=Token, tags=["auth"])
def login(form: OAuth2PasswordRequestForm = Depends(), db: Session = Depends(get_db)):
    user = db.scalar(select(User).where(User.email == form.username.lower()))
    if not user or not verify_password(form.password, user.password_hash):
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "Incorrect email or password")
    return Token(access_token=create_token(user.id))


@app.get("/api/auth/me", tags=["auth"])
def me(user: User = Depends(current_user)):
    return {"id": user.id, "email": user.email}


# ---------- documents ----------

@app.post("/api/documents", response_model=list[UploadResult], status_code=202, tags=["documents"])
async def upload(files: list[UploadFile], user: User = Depends(current_user), db: Session = Depends(get_db)):
    if len(files) > settings.max_files_per_request:
        raise HTTPException(413, f"At most {settings.max_files_per_request} files per request")
    max_bytes = int(settings.max_upload_mb * 1024 * 1024)
    results = []
    for upload_file in files:
        name = (upload_file.filename or "upload.pdf").replace("\\", "/").rsplit("/", 1)[-1][:255]
        data = await upload_file.read(max_bytes + 1)  # never buffer more than the limit
        # Trust the bytes, not the client's filename or Content-Type.
        error = ("File exceeds size limit" if len(data) > max_bytes else
                 "Empty file" if not data else
                 "Only PDF files are accepted" if not data.startswith(b"%PDF-") else None)
        if error:
            results.append(UploadResult(id=None, filename=name, status=None, error=error))
            continue

        sha = hashlib.sha256(data).hexdigest()
        existing = db.scalar(select(Document).where(Document.owner_id == user.id, Document.sha256 == sha))
        if existing:
            results.append(UploadResult(id=existing.id, filename=existing.filename, status=existing.status, duplicate=True))
            continue
        store_file(sha, data)
        doc = Document(owner_id=user.id, filename=name, sha256=sha, size_bytes=len(data), status=Status.queued)
        db.add(doc)
        try:
            db.commit()
        except IntegrityError:  # same file uploaded concurrently; the unique constraint picks a winner
            db.rollback()
            existing = db.scalar(select(Document).where(Document.owner_id == user.id, Document.sha256 == sha))
            results.append(UploadResult(id=existing.id, filename=existing.filename, status=existing.status, duplicate=True))
            continue
        if not enqueue(doc.id):
            doc.error = "Queue unavailable; use Reprocess to retry"
            db.commit()
        db.refresh(doc)
        results.append(UploadResult(id=doc.id, filename=name, status=doc.status))
    return results


@app.get("/api/documents", response_model=list[DocumentSummary], tags=["documents"])
def list_documents(status_filter: Status | None = Query(None, alias="status"),
                   limit: int = Query(100, ge=1, le=500), offset: int = Query(0, ge=0),
                   user: User = Depends(current_user), db: Session = Depends(get_db)):
    q = select(Document).where(Document.owner_id == user.id)
    if status_filter:
        q = q.where(Document.status == status_filter)
    return [summary(d) for d in db.scalars(q.order_by(Document.id.desc()).limit(limit).offset(offset))]


@app.get("/api/documents/{doc_id}", response_model=DocumentDetail, tags=["documents"])
def get_document(doc: Document = Depends(owned_document)):
    return detail(doc)


@app.get("/api/documents/{doc_id}/file", tags=["documents"])
def get_file(doc: Document = Depends(owned_document)):
    return Response(file_path(doc.sha256).read_bytes(), media_type="application/pdf",
                    headers={"Content-Disposition": "inline", "X-Content-Type-Options": "nosniff"})


@app.put("/api/documents/{doc_id}/fields", response_model=DocumentDetail, tags=["review"])
def update_fields(body: Fields, doc: Document = Depends(owned_document), db: Session = Depends(get_db)):
    if doc.status not in (Status.needs_review, Status.validated):
        raise HTTPException(409, f"Cannot edit a document in status {doc.status}")
    doc.fields = body.model_dump(mode="json")
    keep = [i for i in (doc.issues or []) if i["code"] in ("llm_bad_output", "llm_unavailable")]
    doc.issues = keep + validate_document(db, doc, body)
    doc.status, doc.reviewed_at = Status.needs_review, None  # an edit invalidates a previous approval
    db.commit()
    return detail(doc)


@app.post("/api/documents/{doc_id}/approve", response_model=DocumentDetail, tags=["review"])
def approve(doc: Document = Depends(owned_document), db: Session = Depends(get_db)):
    if doc.status != Status.needs_review:
        raise HTTPException(409, f"Cannot approve a document in status {doc.status}")
    issues = validate_document(db, doc, Fields(**doc.fields))  # re-check server-side; never trust the client
    if rules.has_errors(issues):
        raise HTTPException(422, {"message": "Resolve all errors before approving", "issues": issues})
    doc.status, doc.reviewed_at = Status.validated, utcnow()
    db.commit()
    return detail(doc)


@app.post("/api/documents/{doc_id}/reprocess", response_model=DocumentDetail, status_code=202, tags=["documents"])
def reprocess(doc: Document = Depends(owned_document), db: Session = Depends(get_db)):
    if doc.status not in (Status.failed, Status.queued, Status.needs_review):
        raise HTTPException(409, f"Cannot reprocess a document in status {doc.status}")
    doc.status, doc.error = Status.queued, None
    db.commit()
    if not enqueue(doc.id):
        raise HTTPException(503, "Queue unavailable")
    db.refresh(doc)
    return detail(doc)


@app.delete("/api/documents/{doc_id}", status_code=204, tags=["documents"])
def delete_document(doc: Document = Depends(owned_document), db: Session = Depends(get_db)):
    db.delete(doc)
    db.commit()
    # ponytail: the stored file is content-addressed and may be shared; orphan files are left for a sweep job


# ---------- export ----------

HEADER_COLUMNS = ["id", "filename", "document_type", "supplier_name", "customer_name", "document_number",
                  "po_number", "document_date", "due_date", "currency", "subtotal", "tax", "total", "reviewed_at"]
LINE_COLUMNS = ["document_id", "line", "description", "quantity", "unit_price", "amount"]


NUMERIC = {"subtotal", "tax", "total", "quantity", "unit_price", "amount"}


def _cell(col: str, v):
    if v is not None and col in NUMERIC:
        return Decimal(v)
    # Neutralise spreadsheet formula injection from document text (=, +, -, @ at cell start).
    return "'" + v if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r") else v


@app.get("/api/export", tags=["export"])
def export(format: Literal["csv", "xlsx"] = "csv", user: User = Depends(current_user), db: Session = Depends(get_db)):
    docs = db.scalars(select(Document).where(Document.owner_id == user.id, Document.status == Status.validated)
                      .order_by(Document.id)).all()
    rows = [[_cell(c, v) for c, v in zip(HEADER_COLUMNS, [d.id, d.filename, *(d.fields.get(c) for c in HEADER_COLUMNS[2:-1]),
                                                          d.reviewed_at.isoformat() if d.reviewed_at else None])]
            for d in docs]
    stamp = utcnow().strftime("%Y%m%d-%H%M%S")

    if format == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(HEADER_COLUMNS)
        w.writerows(rows)
        return StreamingResponse(iter([buf.getvalue()]), media_type="text/csv",
                                 headers={"Content-Disposition": f'attachment; filename="documents-{stamp}.csv"'})

    wb = Workbook()
    ws = wb.active
    ws.title = "documents"
    ws.append(HEADER_COLUMNS)
    for r in rows:
        ws.append(r)
    lines = wb.create_sheet("line_items")
    lines.append(LINE_COLUMNS)
    for d in docs:
        for i, li in enumerate(d.fields.get("line_items") or [], 1):
            lines.append([_cell(c, v) for c, v in zip(LINE_COLUMNS, [d.id, i, *(li.get(c) for c in LINE_COLUMNS[2:])])])
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="documents-{stamp}.xlsx"'})


@app.get("/api/health", tags=["meta"])
def health():
    return {"status": "ok"}
